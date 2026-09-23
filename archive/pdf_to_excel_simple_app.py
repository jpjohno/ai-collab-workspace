import os
import re
import traceback
from datetime import datetime

import pandas as pd
from openpyxl import Workbook
import pdfplumber

import tkinter as tk
from tkinter import filedialog, messagebox


def sanitize_filename(name: str) -> str:
    name = os.path.splitext(os.path.basename(name))[0]
    name = re.sub(r"[^a-zA-Z0-9 _-]", "", name).strip()
    return name or "output"


def unique_output_path(pdf_path: str, out_dir: str) -> str:
    base = sanitize_filename(pdf_path)
    out = os.path.join(out_dir, f"{base}.xlsx")
    if not os.path.exists(out):
        return out
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(out_dir, f"{base}_{ts}.xlsx")


def write_tables_to_excel(tables, excel_path: str):
    wb = Workbook()
    wb.remove(wb.active)
    for i, (sheet_name, df) in enumerate(tables, start=1):
        ws = wb.create_sheet(title=(sheet_name[:31] if sheet_name else f"Table_{i}"))
        for r_idx, row in enumerate(df.itertuples(index=False), start=1):
            for c_idx, value in enumerate(row, start=1):
                ws.cell(row=r_idx, column=c_idx, value=value)
    wb.save(excel_path)


def extract_tables_pdfplumber(pdf_path: str):
    out_tables = []
    with pdfplumber.open(pdf_path) as pdf:
        for p_idx, page in enumerate(pdf.pages, start=1):
            try:
                tables = page.extract_tables()
            except Exception:
                tables = []
            for t_idx, table in enumerate(tables, start=1):
                df = pd.DataFrame(table).dropna(how="all")
                if df.shape[0] > 0:
                    out_tables.append((f"p{p_idx}_t{t_idx}", df))
    return out_tables


def convert_pdf_to_excel(pdf_path: str, out_dir: str) -> str:
    excel_path = unique_output_path(pdf_path, out_dir)
    tables = extract_tables_pdfplumber(pdf_path)
    if not tables:
        raise RuntimeError("No tables detected (if scanned PDF, OCR is needed).")
    write_tables_to_excel(tables, excel_path)
    return excel_path


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("PDF → Excel Converter")
        self.geometry("520x320")
        self.out_dir = os.path.join(os.path.expanduser("~"), "Desktop")

        tk.Label(self, text="PDF → Excel", font=("Helvetica", 16, "bold")).pack(pady=12)

        tk.Button(self, text="Choose PDF(s)", command=self.choose_files, height=2, width=20).pack(pady=10)
        tk.Button(self, text="Change Output Folder", command=self.choose_out_dir).pack()

        self.status = tk.Label(self, text=f"Output: {self.out_dir}", anchor="w")
        self.status.pack(fill="x", padx=20, pady=10)

        self.log = tk.Text(self, height=10)
        self.log.pack(fill="both", expand=True, padx=20, pady=(0, 15))
        self._log("Ready. Click Choose PDF(s).")

    def _log(self, msg):
        self.log.insert("end", msg + "\n")
        self.log.see("end")

    def choose_out_dir(self):
        d = filedialog.askdirectory(initialdir=self.out_dir)
        if d:
            self.out_dir = d
            self.status.config(text=f"Output: {self.out_dir}")

    def choose_files(self):
        paths = filedialog.askopenfilenames(filetypes=[("PDF files", "*.pdf")])
        if not paths:
            return
        self._log("---")
        self._log(f"Converting {len(paths)} file(s)…")

        ok = 0
        for p in paths:
            try:
                out = convert_pdf_to_excel(p, self.out_dir)
                ok += 1
                self._log(f"✅ {os.path.basename(p)} → {os.path.basename(out)}")
            except Exception as e:
                self._log(f"❌ {os.path.basename(p)} failed: {e}")
                self._log(traceback.format_exc())

        self._log(f"Done. Converted {ok}/{len(paths)}.")

if __name__ == "__main__":
    App().mainloop()
